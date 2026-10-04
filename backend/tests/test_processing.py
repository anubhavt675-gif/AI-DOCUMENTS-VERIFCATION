"""
Phase 2 Test Suite: Document Ingestion, Preprocessing, Normalization,
Fingerprinting, and Standardized Metadata/Structural Analysis.
"""
import io
import os
import tempfile
from pathlib import Path

import docx
import fitz
from PIL import Image
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

# Database override for test isolation
os.environ["DATABASE_URL"] = "sqlite:///./test_temp.db"

from app.database import Base, get_db  # noqa: E402
from app.main import app  # noqa: E402
from app.models.document import Document  # noqa: E402
from app.models.document_analysis import DocumentAnalysis  # noqa: E402
from app.services.document_processing import (
    CorruptedDocumentError,
    DocumentFingerprint,
    EmptyDocumentError,
    PasswordProtectedError,
    UnsupportedFormatError,
    calculate_sha256,
    process_document,
)
from app.services.document_processing.docx_processor import DOCXProcessor
from app.services.document_processing.image_processor import ImageProcessor
from app.services.document_processing.metadata_extractor import (
    parse_date_to_iso,
    standardize_metadata,
)
from app.services.document_processing.pdf_processor import PDFProcessor
from app.services.document_processing.text_cleaner import clean_text, process_text_extraction

SQLALCHEMY_TEST_URL = "sqlite:///./test_temp.db"
test_engine = create_engine(
    SQLALCHEMY_TEST_URL, connect_args={"check_same_thread": False}
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


app.dependency_overrides[get_db] = override_get_db


@pytest.fixture(scope="session", autouse=True)
def setup_test_db():
    Base.metadata.create_all(bind=test_engine)
    yield
    Base.metadata.drop_all(bind=test_engine)


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


# ── In-Memory Sample Generators ──────────────────────────────────────────────

def create_sample_pdf(text: str = "Government Order Ref: 2024-GOV-9876\nConfidential clearance granted.", title: str = "Order 9876", author: str = "Ministry") -> bytes:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((50, 72), text)
    doc.set_metadata({
        "title": title,
        "author": author,
        "subject": "National Security",
        "keywords": "security, classified, order",
    })
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_multipage_pdf() -> bytes:
    doc = fitz.open()
    p1 = doc.new_page()
    p1.insert_text((50, 72), "First Page Section A: Introductory Mandate.")
    p2 = doc.new_page()
    p2.insert_text((50, 72), "Second Page Section B: Operational Guidelines.")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_pdf_with_image() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Document with Embedded Seal.")
    # Create small image
    img = Image.new("RGB", (60, 60), color=(100, 150, 200))
    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    page.insert_image(fitz.Rect(50, 100, 150, 200), stream=img_buf.getvalue())
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.core_properties.title = "Executive Decree 44"
    doc.core_properties.author = "Prime Minister Office"
    doc.core_properties.subject = "Infrastructure Protocol"
    doc.core_properties.keywords = "infrastructure, decree, 2024"

    doc.add_heading("Section 1: General Provisions", level=1)
    doc.add_paragraph("This decree governs the allocation of state computing resources.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Resource"
    table.cell(0, 1).text = "Allocation"
    table.cell(1, 0).text = "Secure Data Center"
    table.cell(1, 1).text = "Node Alpha-7"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_sample_jpeg() -> bytes:
    img = Image.new("RGB", (320, 240), color=(200, 220, 240))
    exif = img.getexif()
    exif[0x010E] = "Security Badge Photo"  # ImageDescription
    exif[0x013B] = "Defense Agency"         # Artist
    buf = io.BytesIO()
    img.save(buf, format="JPEG", exif=exif)
    return buf.getvalue()


def create_sample_png() -> bytes:
    img = Image.new("RGB", (150, 150), color=(240, 240, 240))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# ── Text Cleaning Unit Tests ──────────────────────────────────────────────────

class TestTextCleaning:
    def test_unicode_normalization(self):
        # Ligatures ﬁ and ﬂ should normalize via NFKC to 'fi' and 'fl'
        raw = "The ﬁrst ﬂight was authorized."
        cleaned = clean_text(raw)
        assert "first flight" in cleaned

    def test_control_character_removal(self):
        # Control characters like \x00, \x08 should be stripped; \n preserved
        raw = "Secret\x00Code\x08: 9942\nApproved\x1b."
        cleaned = clean_text(raw)
        assert cleaned == "SecretCode: 9942\nApproved."

    def test_whitespace_and_hyphen_cleanup(self):
        raw = "Gov-   ernment   depart-\nment    approved    this."
        cleaned = clean_text(raw)
        # Note: hyphen followed by newline and lowercase letter resolved to department
        assert "department" in cleaned
        assert "   " not in cleaned

    def test_preserves_numbers_and_punctuation(self):
        raw = "Ref: #SEC-9988; Date: 2024-10-01; Amount: $1,250,000.50; Code: [X9-A]."
        res = process_text_extraction(raw)
        assert "$1,250,000.50" in res.cleaned_text
        assert "2024-10-01" in res.cleaned_text
        assert "#SEC-9988" in res.cleaned_text
        assert res.word_count > 0
        assert res.character_count > 0


# ── Metadata Standardization Unit Tests ──────────────────────────────────────

class TestMetadataStandardization:
    def test_standardizes_empty_fields_to_none(self):
        meta = standardize_metadata({})
        assert meta.title is None
        assert meta.author is None
        assert meta.created_at is None

    def test_date_iso_conversion(self):
        # PDF date format
        pdf_date = "D:20240315123000Z"
        iso = parse_date_to_iso(pdf_date)
        assert iso is not None
        assert "2024-03-15" in iso

        # EXIF date format
        exif_date = "2023:11:05 14:22:10"
        iso_exif = parse_date_to_iso(exif_date)
        assert iso_exif is not None
        assert "2023-11-05" in iso_exif


# ── PDF Processor Unit Tests ──────────────────────────────────────────────────

class TestPDFProcessing:
    def test_extract_normal_text_pdf(self, tmp_path):
        pdf_bytes = create_sample_pdf("Directorate of Defense Protocol 12345.")
        pdf_file = tmp_path / "sample.pdf"
        pdf_file.write_bytes(pdf_bytes)

        fp = DocumentFingerprint("SHA-256", calculate_sha256(pdf_file))
        processor = PDFProcessor()
        res = processor.process(pdf_file, "sample.pdf", "doc-1", fp)

        assert res.page_count == 1
        assert "Directorate of Defense" in res.text.cleaned_text
        assert res.metadata.title == "Order 9876"
        assert res.metadata.author == "Ministry"
        assert res.structure["format"] == "PDF"
        assert res.structure["page_dimensions"][0]["width"] > 0

    def test_extract_multipage_pdf(self, tmp_path):
        pdf_bytes = create_multipage_pdf()
        pdf_file = tmp_path / "multipage.pdf"
        pdf_file.write_bytes(pdf_bytes)

        fp = DocumentFingerprint("SHA-256", calculate_sha256(pdf_file))
        processor = PDFProcessor()
        res = processor.process(pdf_file, "multipage.pdf", "doc-2", fp)

        assert res.page_count == 2
        assert "First Page" in res.text.cleaned_text
        assert "Second Page" in res.text.cleaned_text

    def test_extract_pdf_with_image(self, tmp_path):
        pdf_bytes = create_pdf_with_image()
        pdf_file = tmp_path / "with_img.pdf"
        pdf_file.write_bytes(pdf_bytes)

        fp = DocumentFingerprint("SHA-256", calculate_sha256(pdf_file))
        processor = PDFProcessor()
        res = processor.process(pdf_file, "with_img.pdf", "doc-3", fp)

        assert res.image_count >= 1
        assert res.structure["total_images"] >= 1


# ── DOCX Processor Unit Tests ────────────────────────────────────────────────

class TestDOCXProcessing:
    def test_extract_docx_paragraphs_tables_metadata(self, tmp_path):
        docx_bytes = create_sample_docx()
        docx_file = tmp_path / "sample.docx"
        docx_file.write_bytes(docx_bytes)

        fp = DocumentFingerprint("SHA-256", calculate_sha256(docx_file))
        processor = DOCXProcessor()
        res = processor.process(docx_file, "sample.docx", "doc-docx", fp)

        assert "Section 1: General Provisions" in res.text.cleaned_text
        assert "allocation of state computing resources" in res.text.cleaned_text
        assert "Secure Data Center" in res.text.cleaned_text
        assert res.metadata.title == "Executive Decree 44"
        assert res.metadata.author == "Prime Minister Office"
        assert res.table_count == 1
        assert res.paragraph_count >= 2


# ── Image Processor Unit Tests ───────────────────────────────────────────────

class TestImageProcessing:
    def test_jpeg_with_exif(self, tmp_path):
        jpg_bytes = create_sample_jpeg()
        jpg_file = tmp_path / "badge.jpg"
        jpg_file.write_bytes(jpg_bytes)

        fp = DocumentFingerprint("SHA-256", calculate_sha256(jpg_file))
        processor = ImageProcessor()
        res = processor.process(jpg_file, "badge.jpg", "doc-img-1", fp)

        assert res.structure["width"] == 320
        assert res.structure["height"] == 240
        assert res.structure["has_exif"] is True
        assert res.metadata.title == "Security Badge Photo"
        assert res.ocr.status in ("not_required", "completed", "unavailable", "failed")

    def test_png_without_exif(self, tmp_path):
        png_bytes = create_sample_png()
        png_file = tmp_path / "logo.png"
        png_file.write_bytes(png_bytes)

        fp = DocumentFingerprint("SHA-256", calculate_sha256(png_file))
        processor = ImageProcessor()
        res = processor.process(png_file, "logo.png", "doc-img-2", fp)

        assert res.structure["width"] == 150
        assert res.structure["height"] == 150
        assert res.structure["format"] == "PNG"


# ── Error & Failure Cases ────────────────────────────────────────────────────

class TestFailureCases:
    def test_corrupt_file_raises_corrupted_document_error(self, tmp_path):
        corrupt_file = tmp_path / "corrupt.pdf"
        corrupt_file.write_bytes(b"%PDF-1.4 THIS IS TRUNCATED AND INVALID BYTES")

        with pytest.raises(CorruptedDocumentError):
            process_document(corrupt_file, "corrupt.pdf", "doc-err-1", "application/pdf")

    def test_empty_document_raises_empty_document_error(self, tmp_path):
        empty_file = tmp_path / "empty.pdf"
        empty_file.write_bytes(b"")

        with pytest.raises(EmptyDocumentError):
            process_document(empty_file, "empty.pdf", "doc-err-2", "application/pdf")

    def test_unsupported_format_raises_error(self, tmp_path):
        bad_file = tmp_path / "test.xyz"
        bad_file.write_bytes(b"dummy")

        with pytest.raises(UnsupportedFormatError):
            process_document(bad_file, "test.xyz", "doc-err-3", "application/octet-stream")


# ── Full API Integration Tests (Upload -> Hash -> Process -> Analysis) ───────

class TestAPIIntegration:
    def test_upload_and_analysis_flow(self, client):
        # 1. Upload a PDF
        pdf_bytes = create_sample_pdf("Official Parliamentary Act 2024-SEC-01.")
        file_tuple = ("file", ("parliament_act.pdf", io.BytesIO(pdf_bytes), "application/pdf"))
        up_resp = client.post("/api/documents/upload", files=[file_tuple])
        assert up_resp.status_code == 201
        doc_id = up_resp.json()["document_id"]

        # 2. Trigger preprocessing endpoint
        proc_resp = client.post(f"/api/documents/{doc_id}/process")
        assert proc_resp.status_code == 200
        proc_data = proc_resp.json()
        assert proc_data["success"] is True
        assert proc_data["processing_status"] == "completed"

        # 3. Retrieve analysis result
        ana_resp = client.get(f"/api/documents/{doc_id}/analysis")
        assert ana_resp.status_code == 200
        ana_data = ana_resp.json()

        assert ana_data["document_id"] == doc_id
        assert ana_data["file"]["hash"] is not None
        assert len(ana_data["file"]["hash"]) == 64  # SHA-256 hex length
        assert "Official Parliamentary Act" in ana_data["text"]["cleaned_text"]
        assert ana_data["text"]["word_count"] > 0
        assert ana_data["metadata"]["title"] == "Order 9876"
        assert ana_data["structure"]["format"] == "PDF"

    def test_docx_upload_and_analysis_flow(self, client):
        docx_bytes = create_sample_docx()
        file_tuple = ("file", ("policy.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
        up_resp = client.post("/api/documents/upload", files=[file_tuple])
        assert up_resp.status_code == 201
        doc_id = up_resp.json()["document_id"]

        proc_resp = client.post(f"/api/documents/{doc_id}/process")
        assert proc_resp.status_code == 200
        assert proc_resp.json()["processing_status"] == "completed"

        ana_resp = client.get(f"/api/documents/{doc_id}/analysis")
        assert ana_resp.status_code == 200
        data = ana_resp.json()
        assert data["file"]["hash"] is not None
        assert "General Provisions" in data["text"]["cleaned_text"]
        assert data["structure"]["table_count"] == 1

    def test_corrupt_upload_graceful_failure(self, client):
        # Upload a broken PDF file
        file_tuple = ("file", ("broken.pdf", io.BytesIO(b"%PDF-1.4 CORRUPTED DATA"), "application/pdf"))
        up_resp = client.post("/api/documents/upload", files=[file_tuple])
        assert up_resp.status_code == 201
        doc_id = up_resp.json()["document_id"]

        # Trigger process
        proc_resp = client.post(f"/api/documents/{doc_id}/process")
        assert proc_resp.status_code == 200
        data = proc_resp.json()
        assert data["processing_status"] == "failed"

        # Check analysis endpoint reports failed without crashing
        ana_resp = client.get(f"/api/documents/{doc_id}/analysis")
        assert ana_resp.status_code == 200
        ana_data = ana_resp.json()
        assert ana_data["processing_status"] == "failed"
        assert ana_data["processing_error"] is not None
